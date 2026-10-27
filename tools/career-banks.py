#!/usr/bin/env python3
"""Банки профессий → одно пространство RIASEC, расстояние до человека из candidates.json.

    python3 tools/career-banks.py [fetch]     # fetch — докачать сырьё (нужен openpyxl)

Банки: O*NET 30.0 (интересы есть свои) · ESCO v1.2 (через мост ESCO→O*NET-SOC 2019 с onetcenter.org) ·
ОКЗ ОК 010-2014 (= ISCO-08; группа из 4 цифр получает среднее по профессиям ESCO с тем же кодом) ·
профроли hh.ru (кодов нет — в пространство не кладутся, только список).
Расстояние — как в src/career.js: z-оценки по шкалам O*NET, ось взвешена 1 / (1 + sd²) теста.
Сырьё и результат: папка banks/ в career_data из tools/sources.json.
"""
import csv, json, math, re, sys, time, urllib.request
from sources import ROOT, src

RAW = src('career_data') / 'banks'
UA = {'User-Agent': 'Mozilla/5.0 (X11; Linux x86_64) caerno-research'}
SCALES = 'RIASEC'


def get(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=40) as r:
        return r.read()


def fetch():
    RAW.mkdir(parents=True, exist_ok=True)
    if not (RAW / 'hh-professional-roles.json').exists():
        (RAW / 'hh-professional-roles.json').write_bytes(get('https://api.hh.ru/professional_roles'))
    if not (RAW / 'ESCO_to_ONET-SOC.xlsx').exists():
        (RAW / 'ESCO_to_ONET-SOC.xlsx').write_bytes(get('https://www.onetcenter.org/crosswalks/esco/ESCO_to_ONET-SOC.xlsx'))
    if not (RAW / 'esco-occupations-en.json').exists():
        rows, off = [], 0
        while True:
            d = json.loads(get(f'https://ec.europa.eu/esco/api/search?type=occupation&language=en&limit=100&offset={off}'))
            res = d['_embedded']['results']
            rows += [{'code': r['code'], 'title': r['title'], 'uri': r['uri']} for r in res]
            off += 1                                   # offset у ESCO — номер страницы, не строки
            if len(rows) >= d['total'] or not res:
                break
            time.sleep(0.3)
        (RAW / 'esco-occupations-en.json').write_text(json.dumps(rows, ensure_ascii=False, indent=0))
    if not (RAW / 'okz-ok-010-2014.json').exists():
        names, todo, seen = {}, [str(i) for i in range(10)], set()
        while todo:
            code = todo.pop(0)
            if code in seen or len(code) > 3:
                continue
            seen.add(code)
            html = get(f'https://classifikators.ru/okz/{code}').decode('utf-8', 'replace')
            for c, name in re.findall(r'href="/okz/(\d+)"[^>]*>([^<]+)<', html):
                if c.startswith(code) and len(c) == len(code) + 1 and name.strip():
                    names[c] = name.strip()
                    todo.append(c)
            time.sleep(0.25)
        (RAW / 'okz-ok-010-2014.json').write_text(json.dumps(dict(sorted(names.items())), ensure_ascii=False, indent=0))


def xwalk():
    """{код ESCO: [коды O*NET-SOC]} из листа xlsx."""
    from openpyxl import load_workbook
    out = {}
    for row in load_workbook(RAW / 'ESCO_to_ONET-SOC.xlsx', read_only=True).active.iter_rows(values_only=True):
        if row[0] and row[2] and re.match(r'^\d{2}-\d{4}\.\d{2}$', str(row[2])):
            out.setdefault(str(row[0]), []).append(str(row[2]))
    return out


def main():
    occ = json.loads((ROOT / 'data/career/occupations.json').read_text())
    me = json.loads((ROOT / 'data/career/candidates.json').read_text())['people'][0]
    stats = occ['stats']
    z = {o['code']: [(o['v'][i] - stats[s]['mean']) / stats[s]['sd'] for i, s in enumerate(SCALES)] for o in occ['occupations']}
    w = [1 / (1 + me['sd'][s] ** 2) for s in SCALES]
    dist = lambda v: math.sqrt(sum(w[i] * (me['theta'][s] - v[i]) ** 2 for i, s in enumerate(SCALES)))
    mean = lambda vs: [sum(v[i] for v in vs) / len(vs) for i in range(6)]
    top3 = lambda v: ''.join(s for _, s in sorted(zip(v, SCALES), reverse=True)[:3])

    rows = [{'bank': 'onet', 'code': o['code'], 'title': o['title'], 'n': 1, 'z': z[o['code']]} for o in occ['occupations']]

    esco = json.loads((RAW / 'esco-occupations-en.json').read_text())
    xw = xwalk()
    by_isco, miss = {}, 0
    for e in esco:
        vs = [z[c] for c in xw.get(e['code'], []) if c in z]
        if not vs:
            miss += 1
            continue
        v = mean(vs)
        rows.append({'bank': 'esco', 'code': e['code'], 'title': e['title'], 'n': len(vs), 'z': v})
        by_isco.setdefault(e['code'][:4], []).append(v)

    okz = json.loads((RAW / 'okz-ok-010-2014.json').read_text())
    okz4 = {c: n for c, n in okz.items() if len(c) == 4}
    for c, name in okz4.items():
        if c in by_isco:
            rows.append({'bank': 'okz', 'code': c, 'title': name, 'n': len(by_isco[c]), 'z': mean(by_isco[c])})

    hh = json.loads((RAW / 'hh-professional-roles.json').read_text())
    hh_roles = {r['id']: (r['name'], c['name']) for c in hh['categories'] for r in c['roles']}

    for r in rows:
        r['d'] = round(dist(r['z']), 3)
        r['riasec'] = top3(r['z'])
        r['z'] = [round(x, 2) for x in r['z']]
    for bank in ('onet', 'esco', 'okz'):
        sub = sorted((r for r in rows if r['bank'] == bank), key=lambda r: r['d'])
        for i, r in enumerate(sub, 1):
            r['rank'], r['of'] = i, len(sub)

    with open(RAW / 'banks-distance.csv', 'w', newline='', encoding='utf-8') as f:
        wr = csv.writer(f)
        wr.writerow(['bank', 'rank', 'of', 'd', 'riasec', 'code', 'title', 'n_onet_or_esco', 'zR', 'zI', 'zA', 'zS', 'zE', 'zC'])
        for r in sorted(rows, key=lambda r: (r['bank'], r['rank'])):
            wr.writerow([r['bank'], r['rank'], r['of'], r['d'], r['riasec'], r['code'], r['title'], r['n'], *r['z']])
    with open(RAW / 'hh-roles.csv', 'w', newline='', encoding='utf-8') as f:
        wr = csv.writer(f)
        wr.writerow(['id', 'role', 'category'])
        for i, (n, c) in sorted(hh_roles.items(), key=lambda kv: int(kv[0])):
            wr.writerow([i, n, c])

    print(f"O*NET {len(z)} · ESCO {len(esco)} (без моста/интересов {miss}) · ОКЗ групп-4 {len(okz4)} "
          f"(с ESCO {sum(c in by_isco for c in okz4)}) · hh ролей {len(hh_roles)}")
    print('человек:', me['id'], top3([me['theta'][s] for s in SCALES]), me['source'])


if __name__ == '__main__':
    if sys.argv[1:] == ['fetch']:
        fetch()
    main()
